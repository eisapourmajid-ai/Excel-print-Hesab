"""
Module for safely merging and appending mapped rows into the main Excel system files.
Creates automated backups, locates the first available empty row,
writes input cells, and strictly preserves formula columns and worksheet integrity.
"""

import os
import shutil
from datetime import datetime
from typing import List, Dict, Any
import openpyxl
from app.schemas import TARGET_SCHEMAS


def create_backup(file_path: str) -> str:
    """Creates a timestamped backup in the backups/ directory."""
    os.makedirs("backups", exist_ok=True)
    base_name = os.path.basename(file_path)
    name, ext = os.path.splitext(base_name)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join("backups", f"{name}_{timestamp}{ext}")
    shutil.copy2(file_path, backup_path)
    return backup_path


def find_first_empty_row(ws, start_row: int, max_cap: int, check_cols: List[int]) -> int:
    """Finds the first row where all specified check columns are empty/None."""
    search_limit = max(ws.max_row + 5, max_cap)
    for r in range(start_row, search_limit + 1):
        is_empty = True
        for col_idx in check_cols:
            val = ws.cell(r, col_idx).value
            if val is not None and str(val).strip() != "":
                is_empty = False
                break
        if is_empty:
            return r
    return search_limit + 1


def merge_rows_into_system_excel(
    target_module_id: str,
    rows_to_merge: List[Dict[str, Any]],
    user_name: str = "مدیر سیستم"
) -> Dict[str, Any]:
    """Safely merges rows_to_merge into the specified target Excel file."""
    schema = TARGET_SCHEMAS.get(target_module_id)
    if not schema:
        raise ValueError(f"ماژول نامعتبر: {target_module_id}")
    
    file_name = schema["file"]
    sheet_name = schema["sheet"]
    
    if not os.path.exists(file_name):
        raise FileNotFoundError(f"فایل سیستم {file_name} در مسیر یافت نشد.")
    
    # 1. Create Backup first
    backup_path = create_backup(file_name)
    
    # 2. Load target workbook preserving formulas
    wb = openpyxl.load_workbook(file_name, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"برگه {sheet_name} در فایل {file_name} یافت نشد.")
    
    ws = wb[sheet_name]
    start_row = schema.get("start_row", 9)
    max_cap = schema.get("max_capacity", 2008)
    check_cols = schema.get("check_empty_cols", [2, 4])
    
    # 3. Find first empty row
    current_row = find_first_empty_row(ws, start_row, max_cap, check_cols)
    start_merged_row = current_row
    
    col_map = {}
    for c in schema["columns"]:
        col_map[c["key"]] = {
            "col_index": c["col_index"],
            "type": c.get("type", "text"),
            "title": c["title"]
        }
    
    merged_count = 0
    
    for item in rows_to_merge:
        if item.get("_selected") is False:
            continue
            
        r = current_row
        
        # Write mapped values
        for key, col_info in col_map.items():
            if key in item and item[key] is not None:
                val = item[key]
                col_idx = col_info["col_index"]
                
                # Special formatting for Transactions
                if target_module_id == "transactions":
                    # Col 5: Bank code (ensure bank code is written)
                    if key == "col5" or col_idx == 5:
                        val = item.get("_bank_code", val) or "BANK01"
                    
                    # Col 10: Txn Type (واریز / برداشت)
                    if key == "col10" or col_idx == 10:
                        op_title = str(val)
                        # Determine simple type for Transactions.xlsx (واریز / برداشت)
                        dep_val = float(item.get("col13") or 0)
                        wit_val = float(item.get("col14") or 0)
                        if dep_val > 0 and wit_val == 0:
                            val = "واریز"
                        elif wit_val > 0 and dep_val == 0:
                            val = "برداشت"
                        else:
                            val = "واریز" if "واریز" in op_title else "برداشت"
                    
                    # Col 11: Purpose / Category (Save the detailed accounting operation title!)
                    if key == "col11" or col_idx == 11:
                        # If user selected an operation title in col10, store it in col11 as purpose
                        raw_op = item.get("col10")
                        if raw_op and str(raw_op).strip() != "":
                            val = str(raw_op).strip()
                
                # Format value based on type
                if col_info["type"] == "number":
                    try:
                        if str(val).strip() != "":
                            val = float(str(val).replace(",", "").strip())
                            if val.is_integer():
                                val = int(val)
                    except Exception:
                        pass
                elif col_info["type"] == "date":
                    val = str(val).strip()
                elif col_info["type"] in ["text", "party", "select"]:
                    val = str(val).strip() if str(val).strip() != "" else None
                
                ws.cell(row=r, column=col_idx, value=val)
        
        # For Transactions: Ensure Col 5 (Bank) is set even if not mapped directly
        if target_module_id == "transactions":
            bank_cell = ws.cell(row=r, column=5)
            if bank_cell.value is None or str(bank_cell.value).strip() == "":
                bank_cell.value = item.get("_bank_code", "BANK01")
        
        # Ensure preserved formulas are present
        preserved = schema.get("preserved_formulas", {})
        for f_col, f_template in preserved.items():
            cell = ws.cell(row=r, column=f_col)
            if cell.value is None or str(cell.value).strip() == "":
                r_counter = r - start_row + 1
                formula = f_template.format(r=r, r_counter=r_counter)
                cell.value = formula
        
        merged_count += 1
        current_row += 1
    
    # 4. Save workbook
    wb.save(file_name)
    wb.close()
    
    end_merged_row = current_row - 1
    file_size_kb = round(os.path.getsize(file_name) / 1024, 1)
    
    return {
        "success": True,
        "merged_count": merged_count,
        "target_file": file_name,
        "target_sheet": sheet_name,
        "start_row": start_merged_row,
        "end_row": end_merged_row,
        "backup_path": backup_path,
        "backup_name": os.path.basename(backup_path),
        "file_size_kb": file_size_kb,
        "timestamp": datetime.now().strftime("%Y/%m/%d - %H:%M:%S")
    }
